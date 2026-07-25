"""Generate real multi-format fixtures for the document upload smoke test.

Run with the AI service environment so python-docx / openpyxl / python-pptx
(Docling transitive dependencies) are importable:

    uv run --project services/ai python scripts/generate-smoke-fixtures.py <output-dir>

The PDF is hand-written so the smoke does not need extra dependencies; it uses
Helvetica for Latin text and a standard CID font reference for Chinese text.
"""

from __future__ import annotations

import json
import sys
from pathlib import Path


def pdf_escape(text: str) -> str:
    return text.replace("\\", r"\\").replace("(", r"\(").replace(")", r"\)")


def utf16be_hex(text: str) -> str:
    return text.encode("utf-16-be").hex().upper()


def build_pdf(lines: list[tuple[str, str]]) -> bytes:
    """Build a single-page PDF. lines are (font, text): 'latin' or 'cjk'."""
    content_parts = ["BT", "/F1 14 Tf", "50 780 Td", "20 TL"]
    current_font = "F1"
    for font, text in lines:
        wanted = "F2" if font == "cjk" else "F1"
        if wanted != current_font:
            content_parts.append(f"/{wanted} 14 Tf")
            current_font = wanted
        if font == "cjk":
            content_parts.append(f"<{utf16be_hex(text)}> Tj")
        else:
            content_parts.append(f"({pdf_escape(text)}) Tj")
        content_parts.append("T*")
    content_parts.append("ET")
    content = "\n".join(content_parts).encode("ascii")

    # Identity ToUnicode CMap for the CJK font so text extraction returns the
    # exact codepoints instead of Kangxi radical compatibility characters.
    cjk_codepoints = sorted({ord(char) for font, text in lines if font == "cjk" for char in text})
    cmap_lines = [
        "/CIDInit /ProcSet findresource begin",
        "12 dict begin",
        "begincmap",
        "/CIDSystemInfo << /Registry (Adobe) /Ordering (UCS) /Supplement 0 >> def",
        "/CMapName /Adobe-Identity-UCS def",
        "/CMapType 2 def",
        "1 begincodespacerange",
        "<0000> <FFFF>",
        "endcodespacerange",
    ]
    for start in range(0, len(cjk_codepoints), 100):
        batch = cjk_codepoints[start : start + 100]
        cmap_lines.append(f"{len(batch)} beginbfchar")
        for codepoint in batch:
            cmap_lines.append(f"<{codepoint:04X}> <{codepoint:04X}>")
        cmap_lines.append("endbfchar")
    cmap_lines += [
        "endcmap",
        "CMapName currentdict /CMap defineresource pop",
        "end",
        "end",
    ]
    to_unicode = "\n".join(cmap_lines).encode("ascii")

    objects: list[bytes] = []
    objects.append(b"<< /Type /Catalog /Pages 2 0 R >>")
    objects.append(b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>")
    objects.append(
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 595 842] "
        b"/Resources << /Font << /F1 4 0 R /F2 5 0 R >> >> /Contents 8 0 R >>"
    )
    objects.append(b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")
    objects.append(
        b"<< /Type /Font /Subtype /Type0 /BaseFont /STSong-Light "
        b"/Encoding /UniGB-UCS2-H /DescendantFonts [6 0 R] /ToUnicode 9 0 R >>"
    )
    objects.append(
        b"<< /Type /Font /Subtype /CIDFontType0 /BaseFont /STSong-Light "
        b"/CIDSystemInfo << /Registry (Adobe) /Ordering (GB1) /Supplement 5 >> "
        b"/FontDescriptor 7 0 R >>"
    )
    objects.append(
        b"<< /Type /FontDescriptor /FontName /STSong-Light /Flags 6 "
        b"/FontBBox [-25 -254 1000 880] /ItalicAngle 0 /Ascent 880 "
        b"/Descent -120 /CapHeight 880 /StemV 80 >>"
    )
    objects.append(
        b"<< /Length " + str(len(content)).encode() + b" >>\nstream\n" + content + b"\nendstream"
    )
    objects.append(
        b"<< /Length " + str(len(to_unicode)).encode() + b" >>\nstream\n" + to_unicode + b"\nendstream"
    )

    output = bytearray(b"%PDF-1.4\n")
    offsets = []
    for index, body in enumerate(objects, start=1):
        offsets.append(len(output))
        output += f"{index} 0 obj\n".encode() + body + b"\nendobj\n"
    xref_offset = len(output)
    output += f"xref\n0 {len(objects) + 1}\n".encode()
    output += b"0000000000 65535 f \n"
    for offset in offsets:
        output += f"{offset:010d} 00000 n \n".encode()
    output += (
        f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\n"
        f"startxref\n{xref_offset}\n%%EOF\n"
    ).encode()
    return bytes(output)


def main() -> None:
    output_dir = Path(sys.argv[1] if len(sys.argv) > 1 else ".smoke-logs/fixtures")
    output_dir.mkdir(parents=True, exist_ok=True)

    # --- PDF: English + Chinese -------------------------------------------
    (output_dir / "atlas-quarterly-report.pdf").write_bytes(
        build_pdf(
            [
                ("latin", "Atlas RAG Quarterly Report"),
                ("latin", "Hybrid retrieval combines pgvector and Elasticsearch BM25."),
                ("latin", "Reciprocal rank fusion merges the candidate lists."),
                ("cjk", "知识库平台支持多格式文档导入与可验证引用。"),
                ("cjk", "检索阶段执行严格的访问控制过滤。"),
                ("latin", "Citations open only after a fresh authorization check."),
            ]
        )
    )

    # --- DOCX: Chinese ------------------------------------------------------
    from docx import Document

    document = Document()
    document.add_heading("Atlas 知识库产品需求文档", level=0)
    document.add_heading("目标", level=1)
    document.add_paragraph("为企业提供可私有化部署的检索增强生成平台,支持多格式文档导入。")
    document.add_paragraph("Every answer must carry verifiable citations with fresh ACL checks.")
    document.add_heading("核心能力", level=1)
    for capability in ["混合检索与重排序", "知识图谱治理", "分层记忆", "流式语音交互"]:
        document.add_paragraph(capability, style="List Bullet")
    table = document.add_table(rows=3, cols=2)
    table.rows[0].cells[0].text = "指标"
    table.rows[0].cells[1].text = "目标值"
    table.rows[1].cells[0].text = "召回率 Recall@10"
    table.rows[1].cells[1].text = ">= 0.85"
    table.rows[2].cells[0].text = "引用准确率"
    table.rows[2].cells[1].text = ">= 0.95"
    document.save(output_dir / "产品需求文档.docx")

    # --- XLSX: Chinese, two sheets ------------------------------------------
    from openpyxl import Workbook

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "季度销售"
    sheet.append(["地区", "季度", "销售额(万元)", "同比增长"])
    for row in [
        ("华东", "Q1", 1250.5, "12.3%"),
        ("华北", "Q1", 986.2, "8.1%"),
        ("华南", "Q1", 1102.8, "15.6%"),
        ("West", "Q1", 640.0, "5.4%"),
    ]:
        sheet.append(list(row))
    second = workbook.create_sheet("团队规模")
    second.append(["部门", "人数"])
    second.append(["研发", 42])
    second.append(["Sales", 18])
    workbook.save(output_dir / "季度销售数据.xlsx")

    # --- PPTX: Chinese + English --------------------------------------------
    from pptx import Presentation
    from pptx.util import Inches

    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[0])
    slide.shapes.title.text = "Atlas RAG 项目进展汇报"
    slide.placeholders[1].text = "文档导入闭环已完成"
    second_slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    second_slide.shapes.title.text = "下一步计划"
    body = second_slide.placeholders[1].text_frame
    body.text = "Legacy Office conversion via sandboxed LibreOffice"
    body.add_paragraph().text = "扫描件 OCR 与图片说明"
    body.add_paragraph().text = "Hybrid retrieval publication"
    notes = second_slide.notes_slide.notes_text_frame
    notes.text = "强调原子发布:新版本索引完成前旧版本保持在线。"
    presentation.save(output_dir / "项目进展汇报.pptx")

    # --- TXT / MD / CSV / JSON ----------------------------------------------
    (output_dir / "会议纪要.txt").write_text(
        "Atlas 平台周会纪要\n\n"
        "第一项:文档导入流水线已完成真实验证,覆盖八种格式。\n\n"
        "Second item: the hybrid retrieval milestone starts next week.\n\n"
        "第三项:知识图谱治理需要候选实体审核界面。\n",
        encoding="utf-8",
    )
    (output_dir / "architecture-overview.md").write_text(
        "# Atlas RAG Architecture\n\n"
        "NestJS 是身份、权限与业务数据的唯一入口。\n\n"
        "## Ingestion Pipeline\n\n"
        "1. Browser streams bytes to the quarantine bucket.\n"
        "2. Python worker scans, validates, parses and chunks.\n"
        "3. NestJS publishes the new version atomically.\n\n"
        "## 检索架构\n\n"
        "- 向量检索使用 pgvector\n"
        "- 关键词检索使用 Elasticsearch BM25\n"
        "- 融合策略采用 RRF,重排序输出 Top 10\n",
        encoding="utf-8",
    )
    (output_dir / "员工名录.csv").write_text(
        "姓名,部门,城市,角色\n"
        "张伟,研发,上海,工程师\n"
        "李娜,产品,北京,产品经理\n"
        "Alice Wang,Engineering,Shenzhen,Architect\n"
        "王强,销售,广州,客户经理\n",
        encoding="utf-8",
    )
    (output_dir / "知识库配置.json").write_text(
        json.dumps(
            {
                "space": {"name": "企业知识库", "retrievalMode": "hybrid"},
                "limits": {"maxUploadMb": 200, "并发上传": 3},
                "features": ["引用溯源", "hybrid-search", "知识图谱"],
            },
            ensure_ascii=False,
            indent=2,
        ),
        encoding="utf-8",
    )

    # --- Negative fixtures ----------------------------------------------------
    (output_dir / "corrupted.pdf").write_bytes(
        b"%PDF-1.4\nthis is not a real pdf body, truncated in the middle"
    )
    (output_dir / "disguised.docx").write_text(
        "This is plain text renamed to .docx and must be rejected.", encoding="utf-8"
    )
    # EICAR standard antivirus test string: safe, but every scanner flags it.
    (output_dir / "eicar-test.txt").write_text(
        "X5O!P%@AP[4\\PZX54(P^)7CC)7}$EICAR-STANDARD-ANTIVIRUS-TEST-FILE!$H+H*",
        encoding="ascii",
    )

    for path in sorted(output_dir.iterdir()):
        print(f"{path.name}\t{path.stat().st_size}")


if __name__ == "__main__":
    main()
